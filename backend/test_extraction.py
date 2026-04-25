import os
from pptx import Presentation
import fitz

def extract_text_from_file(file_path: str, ext: str) -> str:
    text = ""
    try:
        if ext == ".pdf":
            doc = fitz.open(file_path)
            for page in doc:
                text += page.get_text() + "\n"
        elif ext in [".ppt", ".pptx"]:
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        text += shape.text + "\n"
    except Exception as e:
        print(f"Extraction error: {e}")
    return text

# Test with a dummy file if possible, or just verify imports
print("Imports successful")
print(f"fitz version: {fitz.__version__}")
