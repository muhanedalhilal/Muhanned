import os
import shutil
import uuid
import PyPDF2
from pptx import Presentation
from fastapi import APIRouter, File, UploadFile, Depends, HTTPException, status, Form
from sqlalchemy.orm import Session
from app.core.security import get_current_user
from app.models.db_user import DBUser
from app.models.course import Course
from app.models.document import Document
from app.models.knowledge_component import KnowledgeComponent
from app.database.database import get_db
from app.core.supabase_client import supabase
from app.core.ai_service import generate_kcs_from_text, generate_kcs_from_file, fallback_kcs_for_resource

router = APIRouter(prefix="/upload", tags=["upload"])

UPLOAD_DIR = "uploads"
os.makedirs(UPLOAD_DIR, exist_ok=True)

ALLOWED_EXTENSIONS = {".pdf", ".ppt", ".pptx"}
MIME_TYPES = {
    ".pdf": "application/pdf",
    ".ppt": "application/vnd.ms-powerpoint",
    ".pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
}

def extract_text_from_file(file_path: str, ext: str) -> str:
    text = ""
    try:
        if ext == ".pdf":
            with open(file_path, 'rb') as f:
                reader = PyPDF2.PdfReader(f)
                for page in reader.pages:
                    text += (page.extract_text() or "") + "\n"
        elif ext in [".ppt", ".pptx"]:
            prs = Presentation(file_path)
            for slide in prs.slides:
                for shape in slide.shapes:
                    if hasattr(shape, "text") and shape.text:
                        text += shape.text + "\n"
    except Exception as e:
        print(f"Extraction error: {e}")
    return text

@router.post("/")
def upload_file(
    file: UploadFile = File(...),
    course_id: int = Form(None),
    current_user: DBUser = Depends(get_current_user),
    db: Session = Depends(get_db)
):
    """
    Accepts educational files securely. 
    1. Saves temp to local disk
    2. Uploads to Supabase Storage
    3. Extracts text
    4. Calls Gemini to generate Knowledge Components (KCs)
    5. Saves Document and KCs strictly to Postgres
    """
    ext = os.path.splitext(file.filename)[1].lower()
    if ext not in ALLOWED_EXTENSIONS:
        raise HTTPException(
            status_code=status.HTTP_400_BAD_REQUEST,
            detail=f"Invalid file type. Only {', '.join(ALLOWED_EXTENSIONS)} are allowed."
        )

    course = None
    if course_id is not None:
        course = (
            db.query(Course)
            .filter(Course.id == course_id, Course.user_id == current_user.id)
            .first()
        )
        if not course:
            raise HTTPException(status_code=404, detail="Course not found")

    safe_filename = f"{uuid.uuid4()}{ext}"
    save_path = os.path.join(UPLOAD_DIR, safe_filename)

    # 1. Save locally first
    try:
        with open(save_path, "wb") as buffer:
            shutil.copyfileobj(file.file, buffer)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Failed to handle temp file: {str(e)}")
    finally:
        file.file.close()

    # 2. Upload to Supabase Storage Bucket ('documents')
    storage_path = f"{current_user.supabase_auth_id}/{safe_filename}"
    try:
        with open(save_path, "rb") as f:
            res = supabase.storage.from_("documents").upload(storage_path, f, {"content-type": file.content_type})
    except Exception as e:
        # Ignore errors if bucket not set up properly or already exists
        print(f"Supabase Upload Warning: {e}")
        # Note: If this fails, we will still process the text locally.

    # 3. Create Document Record in DB
    new_doc = Document(
        filename=file.filename,
        file_type=ext.replace('.', ''),
        supabase_path=storage_path,
        user_id=current_user.id,
        course_id=course_id
    )
    db.add(new_doc)
    db.commit()
    db.refresh(new_doc)

    # 4. Extract text
    extracted_text = extract_text_from_file(save_path, ext)

    # 5. Generate KCs from the uploaded file. Prefer extracted text, then ask
    # Gemini to analyze the original file when extraction returns nothing.
    kcs_data = []
    generated_components = []
    generation_source = "none"
    source_text = extracted_text.strip()
    if source_text:
        kcs_data = generate_kcs_from_text(source_text)
        if kcs_data:
            generation_source = "extracted_text"

    if not isinstance(kcs_data, list) or not kcs_data:
        kcs_data = generate_kcs_from_file(
            save_path,
            mime_type=MIME_TYPES.get(ext) or file.content_type,
            filename=file.filename
        )
        if kcs_data:
            generation_source = "uploaded_file"

    if not isinstance(kcs_data, list) or not kcs_data:
        kcs_data = fallback_kcs_for_resource(None, file.filename)
        if kcs_data:
            generation_source = "file_fallback"

    # 6. Save KCs to Postgres
    new_kcs = []
    for kc in kcs_data:
        if not isinstance(kc, dict):
            continue
        topic = str(kc.get("topic") or "Unknown Topic").strip() or "Unknown Topic"
        content = str(kc.get("content") or "No content provided.").strip() or "No content provided."
        
        new_kc = KnowledgeComponent(
            topic=topic,
            content=content,
            document_id=new_doc.id,
            user_id=current_user.id,
            course_id=course_id
        )
        db.add(new_kc)
        new_kcs.append(new_kc)
    db.commit()

    for new_kc in new_kcs:
        db.refresh(new_kc)
        generated_components.append({
            "id": new_kc.id,
            "text": new_kc.topic,
            "content": new_kc.content,
            "progress": int((new_kc.mastery_prob or 0.1) * 100)
        })

    # Optionally clean up the local temp file to save space
    try:
        os.remove(save_path)
    except:
        pass

    file_url = None
    if new_doc.supabase_path:
        try:
            file_url = supabase.storage.from_("documents").get_public_url(new_doc.supabase_path)
        except Exception:
            pass

    return {
        "message": "File processed and AI KCs generated!",
        "document_id": new_doc.id,
        "kcs_generated": len(generated_components),
        "generation_source": generation_source,
        "extracted_characters": len(source_text),
        "resource": {
            "id": new_doc.id,
            "text": new_doc.filename,
            "type": new_doc.file_type,
            "fileUrl": file_url
        },
        "components": generated_components
    }
